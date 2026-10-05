#!/usr/bin/env python3
"""Export the small derived tables that the figures read from the Galaxy_benchmark run archive.

Usage:
    python analysis/export_galaxy_benchmark_tables.py --source /path/to/Galaxy_benchmark

Reads tracked files of the archive (paulocilasjr/Galaxy_benchmark) and writes, into data/:

- run_scores.csv: one row per scored run of the four primary model configurations (3,816 runs).
- galaxy_traced_runs.csv: Galaxy-condition runs of those configurations with a parsed trace (1,908 runs).
- galaxy_tool_use.csv: one row per traced Galaxy run and installed tool it called, plus one row per run that called
  a user-defined tool (UDT).
- bixbench_failure_causes.csv (Figure 2): primary and secondary cause codes of every incorrect BixBench-Verified-50 run
  from the run-level failure audit (codes only; the audit's answer and note fields are not exported).
- run_execution_errors.csv (Figure 3): failed shell commands and Galaxy jobs in the error state, per run.
- execution_error_types.csv (Figure 3): execution errors by type and by where they occurred, per run.
- compbiobench_task_domains.csv (Figure 3): the domain of each CompBioBench task.
- galaxy_run_steps.csv (Figures 3 and 4): the analysis steps that ran as Galaxy jobs in each traced Galaxy run.
- galaxy_parameter_checks.csv (Figure 3): parameter-check results of installed-tool requests, per benchmark and model.
- galaxy_failure_classes.csv (Figure 3): failed Galaxy requests by failure class, per benchmark.
- replicate_answer_agreement.csv (Figure 4): number of distinct answers among the three runs of each replicate set
  (answers are compared during the export and never written).
- run_tokens_actions.csv (Figure 5): input tokens, cached input tokens and agent actions, per scored run.
- galaxy_interface_calls.csv (Figure 5): requests to the Galaxy interface and characters returned, per traced Galaxy run
  and interface function.
- galaxy_tool_lookup.csv (Figure 5): tools inspected and never run in the same run, per benchmark, from the archive's
  interface-friction table.
- inspectability_counts.csv (Figure 5 text): what each analysis step leaves for inspection after the run, by condition.

The tables hold identifiers, scores, counts, codes and tool identifiers only: no trace text, prompts or answers.
The script prints the archive commit; record it in data/README.md.
"""
import argparse
import json
import re
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
DATA = HERE.parent / "data"
SCORES = "manuscript_narrative/original_layout/analysis/accuracy_primary_runs.csv"
CALLS = "manuscript_narrative/derived/galaxy_calls/calls.csv.gz"
COVERAGE = "manuscript_narrative/derived/galaxy_calls/run_coverage.csv"
ERRORS = "manuscript_material/on_demand/Source_Data_OD_Fig5.xlsx"          # sheet abc_runs, from fig_on_demand.py
LEDGER = "analysis_reports/galaxy_improvement_20260924/v2_trace_friction/ledger.json"
TOKENS = "manuscript_narrative/original_layout/analysis/token_run_observations.csv"
ACTIONS = "manuscript_material/on_demand/Source_Data_OD_Fig6.xlsx"         # sheet abf_runs, from fig_on_demand.py
FRICTION = "manuscript_narrative/original_layout/analysis/token_interface_friction.csv"
FAILURE_CLASSES = "manuscript_narrative/original_layout/analysis/token_failure_classes.csv"
COMPBIO_AUDIT = "CompBio/compBio_overview_audit.json"
# run-level evidence: <folder>/analysis/<task>/history_analysis_evidence.json
EVIDENCE = {"BixBench_50": "BixBench-Verified-50", "CompBio": "CompBioBench", "IWC": "IWC"}
CONDITION = {"Open-ended code condition": "custom code", "Galaxy condition": "Galaxy"}
LEDGER_MODEL = {"GPT-5.5": "GPT-5.5", "Sol": "GPT-5.6 Sol", "Luna": "GPT-5.6 Luna", "DS-Codex": "DeepSeek V4 Pro"}
BENCHMARK = {"BixBench50": "BixBench-Verified-50", "CompBio": "CompBioBench", "IWC": "IWC"}
TRACK = {"open_ended_code": "custom code", "galaxy": "Galaxy"}  # values defined in data/README.md
# Trace model labels of the four primary configurations; the superseded Claude Code harness is left out.
TRACE_MODEL = {
    "codex_gpt_5_5": "GPT-5.5",
    "codex_gpt_5_6_sol": "GPT-5.6 Sol",
    "codex_gpt_5_6_luna": "GPT-5.6 Luna",
    "deepseek_v4_pro_via_codex": "DeepSeek V4 Pro",
    "codex_deepseek_v4_pro_0813": "DeepSeek V4 Pro",
    "codex_deepseek_v4_pro": "DeepSeek V4 Pro",
}
RUN_KEY = ["benchmark", "task_id", "model", "replicate"]
# IWC evidence names three configurations without the codex_ prefix
EVIDENCE_MODEL = {**TRACE_MODEL, "gpt_5_5": "GPT-5.5", "gpt_5_6_sol": "GPT-5.6 Sol", "gpt_5_6_luna": "GPT-5.6 Luna"}


def archive_commit(source: Path) -> str:
    sha = subprocess.run(["git", "-C", str(source), "rev-parse", "--short", "HEAD"], capture_output=True, text=True)
    dirty = subprocess.run(["git", "-C", str(source), "status", "--porcelain", "--", SCORES, CALLS, COVERAGE, ERRORS,
                            LEDGER, TOKENS, ACTIONS, FRICTION, FAILURE_CLASSES, COMPBIO_AUDIT,
                            *[f"{folder}/analysis" for folder in EVIDENCE]],
                           capture_output=True, text=True)
    if dirty.stdout.strip():
        raise SystemExit(f"Source files have uncommitted changes:\n{dirty.stdout}")
    return sha.stdout.strip()


def export_execution_errors(source: Path) -> pd.DataFrame:
    """Failed shell commands (a silent exit code 1 is not counted) and Galaxy jobs in the error state, per run."""
    import openpyxl

    ws = openpyxl.load_workbook(source / ERRORS, read_only=True)["abc_runs"]
    rows = list(ws.iter_rows(values_only=True))
    e = pd.DataFrame(rows[1:], columns=rows[0])
    e = e[e.model_configuration != "DeepSeek V4 Pro (Claude Code, superseded)"]
    out = pd.DataFrame({
        "benchmark": e.benchmark,
        "task_id": e.task,
        "model": e.model_configuration.replace({"DeepSeek V4 Pro (Codex)": "DeepSeek V4 Pro"}),
        "track": e.execution_condition.map(CONDITION),
        "replicate": e.replicate,
        "failed_shell_commands": e.failed_shell_commands.fillna(0).astype(int),
        "galaxy_jobs_in_error_state": e.galaxy_jobs_in_error_state.fillna(0).astype(int),
    }).sort_values(RUN_KEY + ["track"])
    assert out.notna().all().all() and set(out.model) <= set(TRACE_MODEL.values()), "unexpected execution-error rows"
    out.to_csv(DATA / "run_execution_errors.csv", index=False)
    return out


def export_failure_causes(source: Path, scores: pd.DataFrame) -> pd.DataFrame:
    """Primary and secondary cause codes of every incorrect BixBench-Verified-50 run of the primary configurations."""
    led = pd.DataFrame(json.loads((source / LEDGER).read_text()))
    led = led[(led.b == "BixBench") & ~led.run.str.contains("ClaudeCode")]
    out = pd.DataFrame({
        "benchmark": "BixBench-Verified-50",
        "task_id": led.task,
        "model": led.run.str.extract(r"^[CG] (.+) r\d$")[0].map(LEDGER_MODEL),
        "track": led.cond.map(TRACK),
        "replicate": led.run.str.extract(r"r(\d)$")[0].astype(int),
        "primary_cause": led.p,
        "secondary_cause": led.s.replace("-", ""),
        "confidence": led.c,
    }).sort_values(RUN_KEY + ["track"])
    wrong = scores[(scores.benchmark == "BixBench-Verified-50") & (scores.score < 1)]
    key = ["task_id", "model", "track", "replicate"]
    both = wrong.merge(out, on=key, how="outer", indicator=True)
    assert (both._merge == "both").all() and not out.duplicated(key).any(), "every incorrect run needs one cause"
    out.to_csv(DATA / "bixbench_failure_causes.csv", index=False)
    return out


def export_tokens_actions(source: Path, scores: pd.DataFrame) -> pd.DataFrame:
    """Input tokens (cached context included), cached input tokens and agent actions per scored run.

    Actions are the agent's tool calls: shell commands, Galaxy interface calls, web searches or fetches, file reads,
    writes and edits. Runs without a parsed trace have no token or action counts and are left empty."""
    import openpyxl

    t = pd.read_csv(source / TOKENS)
    t = t[t.model_primary]
    t = pd.DataFrame({"benchmark": t.benchmark.map(BENCHMARK), "task_id": t.task, "model": t.cfg,
                      "track": t.env.map(TRACK), "replicate": t.replicate, "input_tokens": t.input_tokens,
                      "cached_input_tokens": t.cached})
    ws = openpyxl.load_workbook(source / ACTIONS, read_only=True)["abf_runs"]
    rows = list(ws.iter_rows(values_only=True))
    a = pd.DataFrame(rows[1:], columns=rows[0])
    a = a[a.model_configuration != "DeepSeek V4 Pro (Claude Code, superseded)"]
    a = pd.DataFrame({"benchmark": a.benchmark, "task_id": a.task,
                      "model": a.model_configuration.replace({"DeepSeek V4 Pro (Codex)": "DeepSeek V4 Pro"}),
                      "track": a.execution_condition.map(CONDITION), "replicate": a.replicate, "actions": a.actions})
    key = RUN_KEY + ["track"]
    assert not t.duplicated(key).any() and not a.duplicated(key).any(), "duplicate token or action rows"
    out = scores[key].merge(t, on=key, how="left").merge(a, on=key, how="left")
    out = out[["benchmark", "task_id", "model", "track", "replicate", "input_tokens", "cached_input_tokens", "actions"]]
    for col in ("input_tokens", "cached_input_tokens", "actions"):
        out[col] = out[col].astype("Int64")
    assert len(out) == len(scores), "token and action rows must match the scored runs"
    out.to_csv(DATA / "run_tokens_actions.csv", index=False)
    return out


def export_interface_calls(source: Path, runs: pd.DataFrame) -> pd.DataFrame:
    """Requests to the Galaxy interface and characters it returned to the agent, per traced Galaxy run and function."""
    c = pd.read_csv(source / CALLS, low_memory=False,
                    usecols=["benchmark", "task", "model", "replicate", "tool", "galaxy_server", "chars_returned"])
    c = c[c.model.isin(TRACE_MODEL) & c.galaxy_server]
    c = c.assign(benchmark=c.benchmark.map(BENCHMARK), task_id=c.task, model=c.model.map(TRACE_MODEL),
                 interface_function=c.tool)
    out = (c.groupby(RUN_KEY + ["interface_function"]).agg(requests=("tool", "size"), chars_returned=("chars_returned", "sum"))
           .reset_index().sort_values(RUN_KEY + ["interface_function"]))
    assert out.merge(runs, on=RUN_KEY, how="left", indicator=True)._merge.eq("both").all(), "calls outside traced runs"
    out.to_csv(DATA / "galaxy_interface_calls.csv", index=False)
    return out


def export_tool_lookup(source: Path) -> pd.DataFrame:
    """Distinct tools an agent inspected in a traced Galaxy run, and how many of them that run never ran, per benchmark."""
    fr = pd.read_csv(source / FRICTION)
    fr = fr[fr.measure == "Inspected tools never run in the same run"]
    out = pd.DataFrame({"benchmark": fr.benchmark.map(BENCHMARK), "tools_inspected": fr.denominator.astype(int),
                        "never_run": fr["count"].astype(int), "percent": fr.percent})
    assert out.notna().all().all() and len(out) == 3, "expected one row per benchmark"
    out.to_csv(DATA / "galaxy_tool_lookup.csv", index=False)
    return out


def read_calls(source: Path, columns: list) -> pd.DataFrame:
    """Galaxy interface calls of the four primary configurations, with display names."""
    c = pd.read_csv(source / CALLS, low_memory=False, usecols=["benchmark", "task", "model", "replicate", "tool",
                                                              "galaxy_server", *columns])
    c = c[c.model.isin(TRACE_MODEL) & c.galaxy_server]
    return c.assign(benchmark=c.benchmark.map(BENCHMARK), task_id=c.task, model=c.model.map(TRACE_MODEL))


def run_model(run: dict) -> str:
    return re.sub(r"_r\d+$", "", re.sub(r"^(galaxy|open_ended_code)_", "", run["run_id"]))


def export_task_domains(source: Path) -> pd.DataFrame:
    tasks = json.loads((source / COMPBIO_AUDIT).read_text())["tasks"]
    out = pd.DataFrame({"task_id": [t["task"] for t in tasks], "domain": [t["domain"] for t in tasks]})
    out = out.sort_values("task_id")
    assert len(out) == 100 and out.domain.notna().all(), "every CompBioBench task needs a domain"
    out.to_csv(DATA / "compbiobench_task_domains.csv", index=False)
    return out


def export_run_steps(source: Path, runs: pd.DataFrame) -> pd.DataFrame:
    """Steps that ran as Galaxy jobs in each traced Galaxy run: each installed tool (Tool Shed identifier without
    version, or built-in identifier) and one step named UDT for any user-defined tool job."""
    c = read_calls(source, ["n_jobs", "tool_id_base"])
    c = c[c.n_jobs.fillna(0) > 0]
    installed = c[c.tool == "run_galaxy_tool_and_wait"].assign(step=lambda x: x.tool_id_base)
    udt = c[c.tool == "run_galaxy_udt_and_wait"].assign(step="UDT")
    out = pd.concat([installed, udt])[RUN_KEY + ["step"]].drop_duplicates().sort_values(RUN_KEY + ["step"])
    assert out.step.notna().all(), "every installed-tool job needs a tool identifier"
    assert out.merge(runs, on=RUN_KEY, how="left", indicator=True)._merge.eq("both").all(), "steps outside traced runs"
    out.to_csv(DATA / "galaxy_run_steps.csv", index=False)
    return out


def export_error_types(source: Path) -> pd.DataFrame:
    """Execution errors by type and by where they occurred, per run (sheet abc_every_error). A Galaxy job error is a UDT
    error when its tool identifier is one the run submitted as a UDT."""
    import openpyxl

    ws = openpyxl.load_workbook(source / ERRORS, read_only=True)["abc_every_error"]
    rows = list(ws.iter_rows(values_only=True))
    e = pd.DataFrame(rows[1:], columns=rows[0])
    e = e[e.model_configuration != "DeepSeek V4 Pro (Claude Code, superseded)"]
    c = read_calls(source, ["tool_id_base", "tool_id_full"])
    udt = c[c.tool == "run_galaxy_udt_and_wait"]
    udt_ids = set(udt.tool_id_base.dropna()) | set(udt.tool_id_full.dropna())
    job = e.channel == "Galaxy job"
    e = e.assign(model=e.model_configuration.replace({"DeepSeek V4 Pro (Codex)": "DeepSeek V4 Pro"}),
                 track=e.execution_condition.map(CONDITION), task_id=e.task, benchmark=e.benchmark.replace(BENCHMARK),
                 channel=np.select([job & e.tool.isin(udt_ids), job], ["Galaxy job, UDT", "Galaxy job, installed tool"],
                                   "shell command"))
    out = (e.groupby(["benchmark", "task_id", "model", "track", "replicate", "channel", "error_type"]).size()
           .rename("errors").reset_index())
    assert out.errors.sum() == len(e) and out.track.notna().all() and out.benchmark.isin(BENCHMARK.values()).all(), \
        "unexpected error rows"
    out.to_csv(DATA / "execution_error_types.csv", index=False)
    return out


def export_parameter_checks(source: Path) -> pd.DataFrame:
    """Parameter-check result of every installed-tool request (prov_status, prov_stage), counted per benchmark and model."""
    c = read_calls(source, ["prov_status", "prov_stage"])
    c = c[c.tool == "run_galaxy_tool_and_wait"]
    c = c.assign(prov_status=c.prov_status.fillna("none"), prov_stage=c.prov_stage.fillna("none"))
    out = c.groupby(["benchmark", "model", "prov_status", "prov_stage"]).size().rename("requests").reset_index()
    out.to_csv(DATA / "galaxy_parameter_checks.csv", index=False)
    return out


def export_failure_classes(source: Path) -> pd.DataFrame:
    t = pd.read_csv(source / FAILURE_CLASSES)
    t = t[t.population == "four_primary_configurations"]
    out = pd.DataFrame({"benchmark": t.benchmark.map(BENCHMARK), "failure_stage": t.failure_stage,
                        "failure_class": t.failure_class, "subclass": t.a1_subclass, "requests": t.calls})
    assert out.benchmark.notna().all(), "unexpected failure-class rows"
    out.to_csv(DATA / "galaxy_failure_classes.csv", index=False)
    return out


def normalize_answer(x):
    """Answers compared as text after trimming and lower-casing; numbers to three significant digits."""
    if x is None:
        return None
    s = re.sub(r"\s+", " ", str(x).strip().strip("\"'").lower())
    try:
        return f"{float(s):.3g}"
    except ValueError:
        return s.replace(" ", "")


def export_answer_agreement(source: Path, scores: pd.DataFrame) -> pd.DataFrame:
    """Number of distinct answers among the three runs of each BixBench-Verified-50 and CompBioBench replicate set; a
    missing answer counts as its own value. Answers are compared in memory and never written."""
    rows = []
    for folder in ("BixBench_50", "CompBio"):
        for path in sorted((source / folder / "analysis").glob("*/history_analysis_evidence.json")):
            d = json.loads(path.read_text())
            for run in d["runs"]:
                if run_model(run) in TRACE_MODEL:
                    rows.append(dict(benchmark=EVIDENCE[folder], task_id=d["task"]["task_id"],
                                     model=TRACE_MODEL[run_model(run)], track=TRACK[run["condition"]],
                                     replicate=run["replicate_id"],
                                     answer=normalize_answer((run.get("outcome") or {}).get("submitted_answer"))))
    a = pd.DataFrame(rows)
    key = RUN_KEY + ["track"]
    assert not a.duplicated(key).any(), "one answer per run"
    s = scores[scores.benchmark != "IWC"][key].merge(a, on=key, how="left")
    out = s.groupby(["benchmark", "task_id", "model", "track"]).answer.agg(
        runs="size", distinct_answers=lambda x: x.fillna("<no answer>").nunique(),
        missing_answers=lambda x: int(x.isna().sum())).reset_index()
    out.to_csv(DATA / "replicate_answer_agreement.csv", index=False)
    return out


def _ids(v) -> int:
    if isinstance(v, list):
        return len(v)
    return len(re.findall(r"'([^']+)'", v)) if isinstance(v, str) else 0


def export_inspectability(source: Path) -> pd.DataFrame:
    """What each analysis step leaves for inspection after the run, counted by condition. Galaxy steps are Galaxy jobs;
    custom-code steps are the agent's shell commands that the run-level evidence labels as analysis."""
    steps, runs = [], []
    for folder in EVIDENCE:
        for path in sorted((source / folder / "analysis").glob("*/history_analysis_evidence.json")):
            for run in json.loads(path.read_text())["runs"]:
                if run_model(run) not in EVIDENCE_MODEL:
                    continue
                track = TRACK[run["condition"]]
                runs.append(dict(track=track,
                                 history=run["evidence_completeness"].get("public_history_contents") == "retrieved"))
                for e in run["events"]:
                    galaxy_step = track == "Galaxy" and e["execution_location"] == "galaxy_job"
                    code_step = (track == "custom code" and e["execution_location"] == "agent_runtime"
                                 and e.get("event_type") == "analysis")
                    if not (galaxy_step or code_step):
                        continue
                    tool = str(e.get("tool") or "")
                    # Tool Shed identifiers end in the tool version; built-in tools take the recorded Galaxy release
                    versioned = galaxy_step and (e.get("tool_version") not in (None, "", "None") or
                                                 tool.startswith("toolshed") and len(tool.split("/")) >= 6 or
                                                 not tool.startswith("toolshed"))
                    steps.append(dict(track=track, command=bool(e.get("command")), tool_version=bool(versioned),
                                      parameters=galaxy_step and e.get("parameters") not in (None, "", "None", "{}"),
                                      outputs=galaxy_step and _ids(e.get("output_artifact_ids")) > 0))
    s, r = pd.DataFrame(steps), pd.DataFrame(runs)
    rows = []
    for track, g in s.groupby("track"):
        for element in ("command", "tool_version", "parameters", "outputs"):
            rows.append(dict(track=track, element=element, unit="analysis steps", with_record=int(g[element].sum()),
                             total=len(g)))
    for track, g in r.groupby("track"):
        rows.append(dict(track=track, element="history_retrieved", unit="runs", with_record=int(g.history.sum()),
                         total=len(g)))
    out = pd.DataFrame(rows)
    out.to_csv(DATA / "inspectability_counts.csv", index=False)
    return out


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--source", required=True, type=Path, help="Path to a Galaxy_benchmark checkout")
    args = p.parse_args()
    print(f"Galaxy_benchmark commit: {archive_commit(args.source)}")

    scores = pd.read_csv(args.source / SCORES)
    scores = pd.DataFrame({
        "benchmark": scores.benchmark.map(BENCHMARK),
        "task_id": scores.task,
        "cluster": scores.cluster,          # bootstrap and permutation unit: BixBench source capsule, otherwise task
        "model": scores.cfg,
        "track": scores.env.map(TRACK),
        "replicate": scores.replicate,
        "score": scores.score,              # 0/1 acceptance; IWC: output agreement 0-1
    })
    assert scores.notna().all().all() and len(scores) == 3816, "unexpected run_scores content"
    scores.to_csv(DATA / "run_scores.csv", index=False)

    cov = pd.read_csv(args.source / COVERAGE)
    cov = cov[cov.model.isin(TRACE_MODEL)]
    runs = pd.DataFrame({"benchmark": cov.benchmark.map(BENCHMARK), "task_id": cov.task,
                         "model": cov.model.map(TRACE_MODEL), "replicate": cov.replicate})
    runs = runs.drop_duplicates().sort_values(RUN_KEY)
    runs.to_csv(DATA / "galaxy_traced_runs.csv", index=False)

    calls = pd.read_csv(args.source / CALLS, low_memory=False,
                        usecols=["benchmark", "task", "model", "replicate", "tool", "galaxy_server", "tool_id_base"])
    calls = calls[calls.model.isin(TRACE_MODEL) & calls.galaxy_server]
    calls = calls.assign(benchmark=calls.benchmark.map(BENCHMARK), task_id=calls.task,
                         model=calls.model.map(TRACE_MODEL))
    installed = calls[calls.tool == "run_galaxy_tool_and_wait"].dropna(subset=["tool_id_base"])
    installed = installed.assign(tool_id=installed.tool_id_base, kind="installed")
    udt = calls[calls.tool == "run_galaxy_udt_and_wait"].assign(tool_id="UDT", kind="udt")
    use = pd.concat([installed, udt])[RUN_KEY + ["tool_id", "kind"]].drop_duplicates().sort_values(RUN_KEY + ["tool_id"])
    assert use.merge(runs, on=RUN_KEY, how="left", indicator=True)._merge.eq("both").all(), "tool use outside traced runs"
    use.to_csv(DATA / "galaxy_tool_use.csv", index=False)

    errors = export_execution_errors(args.source)
    causes = export_failure_causes(args.source, scores)
    print(f"run_execution_errors.csv: {len(errors)} runs; bixbench_failure_causes.csv: {len(causes)} runs")
    tokens = export_tokens_actions(args.source, scores)
    interface = export_interface_calls(args.source, runs)
    lookup = export_tool_lookup(args.source)
    print(f"run_tokens_actions.csv: {len(tokens)} runs ({tokens.input_tokens.notna().sum()} with tokens); "
          f"galaxy_interface_calls.csv: {len(interface)} rows; galaxy_tool_lookup.csv: {len(lookup)} rows")
    print(f"run_scores.csv: {len(scores)} runs; galaxy_traced_runs.csv: {len(runs)} runs; "
          f"galaxy_tool_use.csv: {len(use)} rows")
    new = dict(compbiobench_task_domains=export_task_domains(args.source),
               galaxy_run_steps=export_run_steps(args.source, runs),
               execution_error_types=export_error_types(args.source),
               galaxy_parameter_checks=export_parameter_checks(args.source),
               galaxy_failure_classes=export_failure_classes(args.source),
               replicate_answer_agreement=export_answer_agreement(args.source, scores),
               inspectability_counts=export_inspectability(args.source))
    print("; ".join(f"{name}.csv: {len(t)} rows" for name, t in new.items()))


if __name__ == "__main__":
    main()
